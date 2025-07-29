#!/usr/bin/env python3
import os, re, time, requests, fitz, docx, streamlit as st
from urllib.parse import quote_plus
from collections import Counter
from dotenv import load_dotenv
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
import spacy 
from spacy.cli import download as spacy_download
from pydantic import BaseModel, Field
from crewai import Agent, Task, Crew, Process, LLM
import json

# ─────── 0. Load .env ─────────────────────────────────────────────────────────
load_dotenv()
SPRINGER_KEY_META = os.getenv("SPRINGER_KEY_META")
GROQ_API_KEY      = os.getenv("GROQ_API_KEY")
OPENALEX_BASE     = "https://api.openalex.org/sources"

if not SPRINGER_KEY_META:
    st.error("Please set SPRINGER_KEY_META in your .env")
    st.stop()
if not GROQ_API_KEY:
    st.error("Please set GROQ_API_KEY in your .env")
    st.stop()

# ─────── 1. Load spaCy ─────────────────────────────────────────────────────────
try:
    nlp = spacy.load("en_core_web_sm", disable=["ner"])
except OSError:
    spacy_download("en_core_web_sm")
    nlp = spacy.load("en_core_web_sm", disable=["ner"])

# ─────── 2. Springer session with retries ───────────────────────────────────────
springer = requests.Session()
retry = Retry(
    total=3,
    backoff_factor=1,
    status_forcelist=[429,500,502,503,504],
    allowed_methods=["GET"]
)
adapter = HTTPAdapter(max_retries=retry)
springer.mount("https://", adapter)
springer.mount("http://", adapter)

# ─────── 3. Text‐processing utils ────────────────────────────────────────────────
def extract_abstract(full_text: str) -> str:
    pattern = (
        r"(?mi)^abstract[:\s]*(.*?)"
        r"(?=(?:\n\s*(?:\d+\.\s|Introduction)))"
    )
    m = re.search(pattern, full_text, flags=re.DOTALL)
    return m.group(1).strip() if m else ""

def summarize_abstract(abstract: str) -> str:
    doc = nlp(abstract)
    sents = [s.text.strip() for s in doc.sents]
    return " ".join(sents[:3])

def extract_keywords(text: str, top_k: int = 10) -> list[str]:
    doc = nlp(text)
    cands = [c.text.lower().strip() for c in doc.noun_chunks]
    cands += [t.text.lower() for t in doc if t.pos_ == "PROPN"]
    freq = Counter(cands)
    return [kw for kw,_ in freq.most_common(top_k)]

# ─────── 4. Springer API lookup ─────────────────────────────────────────────────
def springer_journal_search(keywords: list[str], per_keyword_limit: int = 10) -> list[dict]:
    journals, seen = [], set()
    for kw in keywords[:5]:
        params = {
            "q":           kw,
            "p":           1,
            "s":           per_keyword_limit,
            "api_key":     SPRINGER_KEY_META,
            "content-type":"journal",
        }
        try:
            resp = springer.get(
                "https://api.springernature.com/meta/v2/json",
                params=params, timeout=15
            )
            resp.raise_for_status()
        except Exception:
            continue

        count = 0
        for rec in resp.json().get("records", []):
            name = rec.get("publicationName","").strip()
            if not name or name in seen:
                continue
            seen.add(name)

            # collect ISSNs/ISBNs
            issns = rec.get("issn",[]) or []
            for fld in ("printIsbn","electronicIsbn","isbn"):
                v = rec.get(fld)
                if isinstance(v,list): issns += v
                elif isinstance(v,str): issns.append(v)

            journals.append({"name": name, "issns": issns})
            count += 1
            if count >= per_keyword_limit:
                break

        time.sleep(1)

    return journals

# ─────── 5. OpenAlex lookup ──────────────────────────────────────────────────────
def openalex_lookup(name: str, issns: list[str]) -> dict:
    def fetch(filter_q: str) -> list[dict]:
        url = f"{OPENALEX_BASE}?filter={filter_q}&per_page=1"
        r = requests.get(url, timeout=10)
        if not r.ok:
            return []
        return r.json().get("results", [])

    # 1) fuzzy name
    results = fetch(f"display_name.search:{quote_plus(name)}")
    # 2) ISSN fallback
    if not results:
        for issn in issns:
            results = fetch(f"issn:{quote_plus(issn)}")
            if results:
                break

    if not results:
        return {}
    v = results[0]
    return {
        "id":             v.get("id"),
        "display_name":   v.get("display_name"),
        "works_count":    v.get("works_count"),
        "cited_by_count": v.get("cited_by_count"),
        "h_index":        v.get("h_index"),
    }

# ─────── 6. CrewAI analyst setup ─────────────────────────────────────────────────
class Recommendation(BaseModel):
    recommendation: str = Field(..., description="Final venue + justification")

analyst_llm = LLM(model="groq/llama-3.3-70b-versatile", api_key=GROQ_API_KEY)
analyst_agent = Agent(
    role="Analyst",
    backstory=(
        "You have:\n"
        "- a 3‑sentence abstract summary\n"
        "- a list of keywords\n"
        "- Springer journal candidates (name+ISSNs)\n"
        "- their OpenAlex metrics\n"
    ),
    goal="Pick the single best journal + justify by topic‑fit, recency, impact.",
    llm=analyst_llm
)
task_analysis = Task(
    name="analysis",
    description=(
        "Analyze the abstract summary, extracted keywords, Springer journals, "
        "and OpenAlex metrics, then recommend the single best source+journal venue. "
        "Justify your choice by topic fit, recency, and impact metrics." 
        "ABSTRACT SUMMARY:\n{{abstract_summary}}\n\n"
        "KEYWORDS:\n{{keywords}}\n\n"
        "SPRINGER CANDIDATES:\n{{springer_journals}}\n\n"
        "OPENALEX METRICS:\n{{openalex_metrics}}\n\n"
    ),
    agent=analyst_agent,
    output_json=Recommendation,
    expected_output="The best journal venue + justification."
)
crew = Crew(
    agents=[analyst_agent],
    tasks=[task_analysis],
    process=Process.sequential
)


# ... your imports, utils, API functions, crew setup, etc. happen above ...

# ─── Streamlit UI ────────────────────────────────────────────────────────
st.set_page_config(page_title="Find My Venue", layout="wide")
st.title("📑 Find My Venue")

# 1) Input selection
st.subheader("📥 Input your paper")
col1, col2 = st.columns(2)
with col1:
    uploaded = st.file_uploader("Upload PDF or DOCX", type=["pdf", "docx"])
with col2:
    manual_text = st.text_area("Or paste your paper text here", height=200)

if not uploaded and not manual_text.strip():
    st.info("Please either upload a PDF/DOCX or paste your text above.")
    st.stop()

# 2) One-shot Analyze button
if st.button("🔎 Analyze Paper"):
    # reset existing placeholders
    prog = st.progress(0)
    spinner = st.empty()

    # 2.1) Full text extraction
    spinner.text("📖 Extracting full text…")
    if uploaded and uploaded.type == "application/pdf":
        import fitz
        doc = fitz.open(stream=uploaded.read(), filetype="pdf")
        full_text = "\n".join(page.get_text() for page in doc)
    elif uploaded:  # docx
        import docx
        doc = docx.Document(uploaded)
        full_text = "\n".join(p.text for p in doc.paragraphs)
    else:
        full_text = manual_text
    prog.progress(10)

    # 2.2) Abstract
    spinner.text("✂️ Extracting abstract…")
    abstract = extract_abstract(full_text)
    if not abstract:
        spinner.empty()
        st.error("❌ Could not locate an Abstract section.")
        st.stop()
    prog.progress(25)

    # 2.3) Summarize
    spinner.text("📝 Summarizing abstract…")
    summary = summarize_abstract(abstract)
    prog.progress(40)

    # 2.4) Keywords
    spinner.text("🔑 Extracting keywords…")
    keywords = extract_keywords(summary)
    prog.progress(55)

    # 2.5) Springer
    spinner.text("📚 Fetching Springer journal candidates…")
    springer_journals = springer_journal_search(keywords, per_keyword_limit=10)
    prog.progress(70)

    # 2.6) OpenAlex
    spinner.text("📊 Looking up OpenAlex metrics…")
    openalex_metrics  = [openalex_lookup(j["name"], j["issns"]) for j in springer_journals]
    prog.progress(85)

    # 2.7) Final recommendation
    spinner.text("🏆 Running final analysis…")
    # you could use return_direct=True here if you like
    result = crew.kickoff(inputs={
        "abstract_summary":  summary,
        "keywords":          keywords,
        "springer_journals": springer_journals,
        "openalex_metrics":  openalex_metrics
    })
    # pull out the string only
    rec = json.loads(result.raw)["recommendation"]
    prog.progress(100)
    spinner.empty()

    # 3) Display all results
    st.markdown("---")
    st.subheader("📝 Extracted Abstract")
    st.write(abstract)

    st.subheader("✂️ 3‑Sentence Summary")
    st.write(summary)

    st.subheader("🔑 Keywords")
    st.write(keywords)

    st.subheader("📚 Springer Journal Candidates")
    if springer_journals:
        st.table(springer_journals)
    else:
        st.write("No matches found.")

    st.subheader("📊 OpenAlex Venue Metrics")
    if any(openalex_metrics):
        st.table(openalex_metrics)
    else:
        st.write("No metrics found.")

    st.subheader("🏆 Final Recommendation")
    st.success(rec)

    # done
